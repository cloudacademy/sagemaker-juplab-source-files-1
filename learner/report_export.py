"""Export edited Markdown report templates to Word and a submission ZIP within Jupyter."""
from pathlib import Path
import re,zipfile
from docx import Document
from docx.shared import Pt,Inches
from labkit import ROOT

def markdown_to_docx(source,target):
    doc=Document();doc.styles['Normal'].font.name='Calibri';doc.styles['Normal'].font.size=Pt(10)
    for section in doc.sections:section.left_margin=section.right_margin=Inches(.7)
    lines=source.read_text().splitlines();i=0
    clean=lambda s:s.replace('**','').replace('`','')
    while i<len(lines):
        line=lines[i]
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                values=[clean(x.strip()) for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(':?-+:?',v) for v in values):rows.append(values)
                i+=1
            if rows:
                table=doc.add_table(rows=0,cols=len(rows[0]));table.style='Light Shading Accent 1'
                for row in rows:
                    cells=table.add_row().cells
                    for j,value in enumerate(row[:len(cells)]):cells[j].text=value
            continue
        if line.startswith('#'):
            level=len(line)-len(line.lstrip('#'));doc.add_heading(clean(line.lstrip('#').strip()),level=min(level,3))
        elif line.startswith('- '):doc.add_paragraph(clean(line[2:]),style='List Bullet')
        elif line.strip():doc.add_paragraph(clean(line))
        i+=1
    doc.save(target)

def export_all():
    base=ROOT;dest=base/'submissions';dest.mkdir(exist_ok=True)
    for p in sorted((base/'templates').glob('Task*.md')):
        if '[learner ID]' in p.read_text():print('WARNING: unfinished placeholder in',p.name)
        markdown_to_docx(p,dest/p.with_suffix('.docx').name)
    archive=dest/'E179_submission.zip'
    with zipfile.ZipFile(archive,'w',zipfile.ZIP_DEFLATED) as z:
        for folder in ['templates','notebooks','outputs']:
            for p in sorted((base/folder).rglob('*')):
                if p.is_file() and '__pycache__' not in p.parts and '.ipynb_checkpoints' not in p.parts:
                    z.write(p,p.relative_to(base))
        for p in sorted(dest.glob('*.docx')):z.write(p,'reports/'+p.name)
        for rel in ['data/manifest.json','models/provenance.json','environment.json']:
            p=base/rel
            if p.exists():z.write(p,'provenance/'+p.name)
    print('Save notebooks before exporting. Review reports and evidence, then download:',archive)
    return archive
if __name__=='__main__':export_all()
